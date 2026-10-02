"""Gera docs/Relatorio_Sprint5.docx lendo os resultados REAIS de evidencias/resultados_experimento.json."""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
EV = RAIZ / "evidencias"
R = json.loads((EV / "resultados_experimento.json").read_text(encoding="utf-8"))
NOMES = {"regras": "Regras puras", "ml": "PLN isolado", "hibrido": "Híbrido (proposto)"}

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(1)
sec.top_margin = sec.bottom_margin = Inches(0.9)
st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(11)
for nome, tam in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    h = doc.styles[nome]; h.font.name = "Calibri"; h.font.size = Pt(tam); h.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)


def p(texto, negrito=False, italico=False, centro=False, espaco=6):
    par = doc.add_paragraph()
    run = par.add_run(texto); run.bold = negrito; run.italic = italico
    par.paragraph_format.space_after = Pt(espaco)
    if centro:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return par


def bullets(itens):
    for i in itens:
        par = doc.add_paragraph(style="List Bullet"); par.add_run(i); par.paragraph_format.space_after = Pt(2)


def sombrear(cell, cor):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), cor); tcPr.append(shd)


def tabela(cab, linhas, larguras=None):
    t = doc.add_table(rows=1, cols=len(cab)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cab):
        cel = t.rows[0].cells[i]; cel.text = ""; r = cel.paragraphs[0].add_run(c); r.bold = True; r.font.size = Pt(10)
        sombrear(cel, "DBEAFE")
    for lin in linhas:
        cels = t.add_row().cells
        for i, v in enumerate(lin):
            cels[i].text = ""; r = cels[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(10)
    if larguras:
        for row in t.rows:
            for i, w in enumerate(larguras):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figura(arquivo, legenda, largura=6.0):
    doc.add_picture(str(EV / arquivo), width=Inches(largura))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = p(legenda, italico=True, centro=True); c.runs[0].font.size = Pt(9.5)


# ----------------------------------------------------------------- capa
p("UNAMA — Universidade da Amazônia", negrito=True, centro=True, espaco=2)
p("Análise e Desenvolvimento de Sistemas", centro=True, espaco=18)
t = doc.add_heading("Sprint 5 — Execução", 0); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Arquitetura Híbrida de Baixo Custo para Triagem Automatizada e Classificação de Chamados Acadêmicos",
  negrito=True, centro=True, espaco=12)
p("Aluno: Caio Vitor Pinheiro Moraes — Matrícula: 04179424", centro=True, espaco=2)
p("Parte Acadêmica (evidências do experimento) e Parte Profissional (MVP Alpha)", centro=True, espaco=14)

# ----------------------------------------------------------------- acadêmico
doc.add_heading("Parte Acadêmica — Execução da metodologia", 1)
doc.add_heading("1. O que foi executado", 2)
p("Seguindo a metodologia definida nas Sprints 2 a 4, a arquitetura híbrida foi implementada e submetida ao "
  "experimento planejado. Foram executadas as etapas: (i) construção e rotulagem do dataset; (ii) implementação dos três "
  "classificadores comparados (regras puras, PLN isolado e híbrido); (iii) validação cruzada com métricas de "
  "recuperação de informação; (iv) teste de robustez a erros de digitação; (v) medição de latência em processo e sob carga HTTP.")
bullets([
    "Dataset: 135 chamados acadêmicos sintéticos (45 por setor: Financeiro, Secretaria, Coordenação), escritos à mão com "
    "variação de linguagem, rotulados com setor correto e urgência (alta/média/baixa) — arquivo data/dataset.csv.",
    "Regras: léxico ponderado por setor (termos fortes e fracos) sobre texto normalizado; decide sozinho quando um setor "
    "vence por margem ≥ 2 pontos.",
    "PLN leve: TF-IDF (palavras 1-2 gramas + n-gramas de caracteres 2-5) com Regressão Logística; acionado quando as regras "
    "são incertas. Chamados com confiança < 0,5 são sinalizados para revisão manual.",
    "Protocolo: validação cruzada estratificada 5-fold repetida 5 vezes (25 avaliações por método). O classificador de PLN é "
    "treinado apenas nas dobras de treino. Cenário B aplica ruído de digitação (troca/omissão/duplicação de letras em ~30% "
    "das palavras) apenas aos textos de teste.",
])

doc.add_heading("2. Resultados", 2)
doc.add_heading("2.1 Acurácia no direcionamento de setor", 3)
linhas = []
for m in ("regras", "ml", "hibrido"):
    a, b = R["acuracia"]["limpo"][m], R["acuracia"]["ruido"][m]
    linhas.append([NOMES[m], f'{a["media"]:.1f}% ± {a["desvio"]:.1f}', f'{b["media"]:.1f}% ± {b["desvio"]:.1f}'])
tabela(["Método", "Texto limpo (média ± dp)", "Com erros de digitação (média ± dp)"], linhas, [2.2, 2.1, 2.3])
figura("fig_acuracia.png", "Figura 1 — Acurácia por método e cenário (validação cruzada 5x5). Linha tracejada: meta de 85%.", 5.4)

doc.add_heading("2.2 Precisão, revocação e F1 por setor (algoritmo híbrido)", 3)
linhas = []
for s, v in R["metricas_por_setor_hibrido"]["limpo"].items():
    n = R["metricas_por_setor_hibrido"]["ruido"][s]
    linhas.append([s, f'{v["precisao"]:.3f}', f'{v["recall"]:.3f}', f'{v["f1"]:.3f}',
                   f'{n["precisao"]:.3f}', f'{n["recall"]:.3f}', f'{n["f1"]:.3f}'])
tabela(["Setor", "P (limpo)", "R (limpo)", "F1 (limpo)", "P (ruído)", "R (ruído)", "F1 (ruído)"], linhas,
       [1.3, 0.9, 0.9, 0.9, 0.9, 0.9, 0.9])
figura("fig_matriz_confusao.png", "Figura 2 — Matriz de confusão do algoritmo híbrido (média por repetição; 45 chamados por setor por repetição).", 5.8)
cob = R["cobertura_regras_no_hibrido"]; urg = R["acuracia_urgencia_regras"]
p(f'Cobertura das regras no híbrido: {cob["limpo"]}% dos chamados foram resolvidos só pelas regras em texto limpo e '
  f'{cob["ruido"]}% com ruído (o restante acionou o PLN). Acurácia da classificação de urgência (regras): '
  f'{urg["media"]}% ± {urg["desvio"]}.')

doc.add_heading("2.3 Desempenho (latência)", 3)
lp = R["latencia_processo"]
tabela(["Método", "Latência média (ms)", "p95 (ms)", "Amostras"],
       [[NOMES[m], lp[m]["media_ms"], lp[m]["p95_ms"], lp[m]["n"]] for m in ("regras", "ml", "hibrido")], [2.2, 1.7, 1.4, 1.3])
lh = R["latencia_http"]
tabela(["Requisições simultâneas", "Total de requisições", "Sucesso", "p50 (ms)", "p95 (ms)", "Vazão (req/s)"],
       [[c, v["requisicoes"], f'{v["sucesso"]}/{v["requisicoes"]}', v["p50_ms"], v["p95_ms"], v["req_por_s"]] for c, v in lh.items()],
       [1.4, 1.2, 0.9, 0.9, 0.9, 1.1])
figura("fig_latencia_http.png", "Figura 3 — Latência de ponta a ponta de POST /api/tickets (triagem + gravação no banco) sob concorrência.", 4.6)
p("Medição feita em máquina única, com o servidor de desenvolvimento do Flask e SQLite; os valores servem como ordem de "
  "grandeza e devem ser repetidos no ambiente de implantação.", italico=True)

doc.add_heading("3. Análise dos resultados", 2)
hl, hr = R["acuracia"]["limpo"]["hibrido"]["media"], R["acuracia"]["ruido"]["hibrido"]["media"]
bullets([
    f"A meta de acurácia ≥ 85% foi atingida pelo algoritmo híbrido nos dois cenários ({hl}% em texto limpo e {hr}% com ruído), "
    "embora, no cenário com ruído, a margem seja estreita em relação ao desvio-padrão entre dobras.",
    "Em texto limpo o híbrido não superou as regras puras (92,0% vs 92,6%, diferença dentro do desvio-padrão). A vantagem do "
    "híbrido aparece com erros de digitação: 86,1% contra 76,7% tanto das regras quanto do PLN isolado (+9,4 p.p.), o que "
    "sustenta a hipótese central: regras dão previsibilidade e custo quase nulo; o PLN cobre os casos em que as regras falham.",
    "O PLN isolado teve o pior desempenho em texto limpo (80,6%), coerente com a literatura: exige base histórica maior do que "
    "as ~108 amostras de treino por dobra usadas aqui.",
    "O custo computacional é baixo: triagem média de 0,34 ms no híbrido (0,04 ms só regras; 1,5 ms PLN), e 100% de sucesso nas "
    "requisições HTTP com até 50 simultâneas, com aumento de latência (p95) conforme a concorrência cresce.",
    "A classificação de urgência por regras (74,1%) é o ponto mais fraco e é candidata a melhoria (classificador próprio).",
])

doc.add_heading("4. Ameaças à validade e limitações (declaradas)", 2)
bullets([
    "Dataset sintético e pequeno (135 chamados): os resultados indicam viabilidade, não desempenho em produção. "
    "O léxico de regras e o dataset foram elaborados por uma mesma origem de conhecimento, o que tende a favorecer as regras "
    "(otimismo). Para a versão final do TCC recomenda-se acrescentar chamados reais anonimizados e/ou textos escritos por "
    "terceiros e repetir o experimento.",
    "O ruído de digitação é simulado por um gerador simples; ele não cobre abreviações, gírias ou textos longos.",
    "O baseline manual (comparação com triagem humana) ainda NÃO foi medido: o protocolo e o script estão prontos "
    "(experimentos/baseline_manual.py), mas exigem voluntários reais. Sem esses dados, a afirmação de “redução drástica de "
    "tempo em relação ao manual” não pode ser feita com evidência — apenas o tempo do sistema (milissegundos) foi medido.",
    "Latência medida em ambiente de desenvolvimento, sem rede real entre cliente e servidor.",
])

doc.add_heading("5. Reprodutibilidade", 2)
p("Todos os números acima foram gerados por python -m experimentos.experimento (sementes fixas) e estão em "
  "evidencias/resultados_experimento.json; o log completo está em evidencias/log_experimento.txt.")

# ----------------------------------------------------------------- profissional
doc.add_page_break()
doc.add_heading("Parte Profissional — MVP Alpha", 1)
doc.add_heading("1. O que funciona", 2)
p("O MVP Alpha é a primeira versão executável do sistema: o fluxo central (aluno abre chamado → API tria → chamado "
  "entra na fila do setor → atendente responde) funciona de ponta a ponta, com interface web.")
tabela(["Requisito do MVP (Sprint 3)", "Situação no Alpha", "Evidência"], [
    ["Autenticação e perfis (Aluno / Atendente) com JWT", "Funcionando", "Teste 01-03, 06; tela de login"],
    ["Gestão de tickets (criar e acompanhar)", "Funcionando", "Teste 04; Figura 4"],
    ["API de triagem automática (setor + urgência)", "Funcionando", "Teste 04, 07; Figuras 1-3"],
    ["Painel do atendente (fila por setor, responder, status)", "Funcionando", "Teste 04-05; Figuras 5-6"],
    ["Reencaminhamento manual de setor (correção da triagem)", "Funcionando (extra)", "Teste 05"],
    ["Notificação por e-mail / busca por palavra-chave / métricas avançadas", "Não iniciado (desejáveis)", "—"],
], [3.3, 1.6, 1.7])

doc.add_heading("2. Arquitetura e stack", 2)
bullets([
    "Backend: Python 3.12 + Flask (API REST). Autenticação: JWT (PyJWT, HS256, expiração de 8 h) e senhas com hash (Werkzeug).",
    "Banco de dados: SQLite (3 tabelas: users, tickets, respostas). Containerização: Dockerfile com gunicorn (ainda não executado neste ambiente).",
    "Frontend: HTML5 + CSS3 + JavaScript puro em página única, servida pela própria API.",
    "Motor de triagem: módulo app/triagem.py (regras + scikit-learn), treinado no início da aplicação a partir de data/dataset.csv.",
    "Desvio em relação ao planejado: as Sprints 2 e 3 previam Java/Spring Boot + PostgreSQL. Para garantir um protótipo "
    "executável e testado nesta etapa, o Alpha foi feito em Python/Flask + SQLite. A arquitetura (API REST, JWT, banco "
    "relacional, contêiner, triagem híbrida) é equivalente e a migração para Spring Boot/PostgreSQL permanece no plano.",
])

doc.add_heading("3. Evidências de funcionamento", 2)
p("Testes automatizados: 8 testes de integração (evidencias/testes_api.txt) passaram, cobrindo login inválido, rotas protegidas, "
  "cadastro e e-mail duplicado, fluxo completo de triagem e painel, isolamento entre setores e entre alunos, reencaminhamento, "
  "validações e tolerância a erros de digitação.")
figura("tela_02_aluno_chamado_triado.png", "Figura 4 — Aluno abre chamados; o sistema informa setor, urgência, método e tempo de triagem.", 5.6)
figura("tela_03_painel_atendente.png", "Figura 5 — Painel do atendente com indicadores e fila ordenada por urgência.", 5.6)
figura("tela_04_atendente_responde.png", "Figura 6 — Detalhe do chamado: resposta do atendente e mudança automática de status para “em andamento”.", 5.6)

doc.add_heading("4. Como executar (demonstração)", 2)
bullets([
    "pip install -r requirements.txt  →  python -m app.app  →  abrir http://localhost:5000",
    "Aluno: aluno@unama.test / aluno123. Atendentes: financeiro@, secretaria@, coordenacao@ ou supervisor@unama.test / atende123.",
])

doc.add_heading("5. Próximos passos", 2)
bullets([
    "Executar o baseline manual (30 chamados, 2-3 avaliadores) e incluir a comparação no relatório final.",
    "Ampliar e diversificar o dataset (chamados reais anonimizados) e repetir o experimento.",
    "Melhorar a classificação de urgência; implementar notificações por e-mail e busca por palavra-chave.",
    "Migrar para Spring Boot + PostgreSQL e publicar (Docker/Render) conforme o plano de portfólio.",
])

doc.save(RAIZ / "docs" / "Relatorio_Sprint5.docx")
print("Relatório gerado.")
